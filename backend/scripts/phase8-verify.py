"""Read-only, independent verification of the synthetic export files."""
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[2] / "outputs" / "phase8-check"
for count in (0, 3, 105):
    report = json.loads((root / f"synthetic-{count}.json").read_text(encoding="utf-8"))
    book = load_workbook(root / f"synthetic-{count}.xlsx")
    assert book.sheetnames == ["Summary", "Performance", "Payments", "Jobs", "Comparison", "Tunnels", "Congestion", "Liabilities"]
    jobs = book["Jobs"]
    rows = list(jobs.iter_rows(min_row=5, values_only=True))
    ids = [row[0] for row in rows if isinstance(row[0], str) and row[0].startswith("SAMPLE-")]
    assert len(ids) == count and len(set(ids)) == count
    headers = {cell.value: cell.column for cell in jobs[4]}
    if count:
        money = headers["Finance / Expected GBP"]
        actual = sum(jobs.cell(index, money).value for index in range(5, count + 5))
        assert round(actual * 100) == report["summary"]["expectedKnownPence"]
        assert jobs.cell(5, money).data_type == "n"
        liability_headers = {cell.value: cell.column for cell in book["Liabilities"][4]}
        assert book["Liabilities"].cell(5, liability_headers["Description"]).value == report["liabilities"][0]["description"]
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    assert all(sheet.freeze_panes == "A5" and sheet.auto_filter.ref for sheet in book)
    pdf = PdfReader(root / f"synthetic-{count}.pdf")
    pages = [page.extract_text() for page in pdf.pages]
    assert all("TMV | DRIVER REPORTS & PERFORMANCE" in text for text in pages), "Unexpected blank/extra footer-only page"
    text = "\n".join(pages)
    # Built-in PDF kerning may be extracted with whitespace before hyphens.
    normalized = re.sub(r"\s+-", "-", text)
    for job in report["jobs"]:
        assert job["jobId"] in normalized
    if count:
        for source in ("EVENT-1", "VISIT-1", "CHARGE-1", "SOURCE-1"):
            assert source in normalized
        assert text.count("Sample evidence details.") == 150
        assert "System calculation" in text and "Driver manual details" in text
    assert f"Page {len(pages)}" in pages[-1]
    for index, page in enumerate(pages):
        assert f"| Page {index + 1}" in page
    print(f"PASS: {count} jobs; eight sheets; matching totals; no formulas; {len(pages)} PDF pages; all identifiers/long description retained.")
